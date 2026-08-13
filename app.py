from flask import Flask,jsonify,request,render_template,session,redirect,render_template_string,send_file,abort
import os,re,math
from io import BytesIO
from collections import Counter
from module2_routes import module2
from module2_engine import QUESTION_BANK,evaluate_answer,select_next,score_assessment
from module3_routes import module3
from module4_routes import module4
from module5_routes import module5
from auth_routes import auth,current_user,require_auth
from auth_db import init_db,record_performance,list_user_summaries,get_user_performance,get_user_activity

app=Flask(__name__,static_folder='static',template_folder='templates')
app.secret_key=os.getenv('FLASK_SECRET_KEY',os.getenv('SESSION_SECRET','dev-only-change-me'))
app.config.update(SESSION_COOKIE_HTTPONLY=True,SESSION_COOKIE_SECURE=True,SESSION_COOKIE_SAMESITE='Lax')
app.register_blueprint(auth);app.register_blueprint(module2);app.register_blueprint(module3);app.register_blueprint(module4);app.register_blueprint(module5)

SKILLS=['python','javascript','react','node.js','java','sql','mongodb','postgresql','rest api','docker','kubernetes','aws','fastapi','flask','django','pytorch','tensorflow','nlp','machine learning','deep learning','scikit-learn','system design','microservices','git','linux','pandas','numpy']
DEMO={'name':'Alex Johnson','role':'Software Engineer','company':'TechNova','skills':['Python','React','JavaScript','Node.js','MongoDB','REST API','Git','Pandas','NumPy'],'experience':3}
JOB={'company':'TechNova','role':'Software Engineer','required':['Python','React','Node.js','REST API','Git'],'preferred':['Docker','AWS','System Design']}

def toks(s):return re.findall(r'[a-z0-9+#./-]+',str(s).lower())
def skillset(s):
 t=' '.join(toks(s));return[x for x in SKILLS if re.search(r'(?<![a-z0-9])'+re.escape(x)+r'(?![a-z0-9])',t)]
def cosine(a,b):
 k=set(a)|set(b);d=sum(a.get(x,0)*b.get(x,0) for x in k);na=math.sqrt(sum(v*v for v in a.values()));nb=math.sqrt(sum(v*v for v in b.values()));return d/(na*nb) if na and nb else 0
def analyze(p):
 r=p.get('resume',DEMO);j=p.get('job',JOB);rt=' '.join(map(str,r.values()));jt=' '.join(map(str,j.values()));rs=set(skillset(rt));req=skillset(jt);found=[x for x in req if x in rs];missing=[x for x in req if x not in rs];rv=Counter(toks(rt));jv=Counter(toks(jt));sim=cosine(rv,jv)*100;sm=len(found)/max(len(req),1)*100;kw=len(set(toks(jt))&set(toks(rt)))/max(len(set(toks(jt))),1)*100;ats=round(.3*kw+.3*sm+.2*sim+.2*95);return{'atsScore':ats,'overallMatch':round((ats+sm+sim)/3),'skillsMatch':round(sm),'keywordMatch':round(kw),'semanticSimilarity':round(sim),'missingSkills':missing,'shortlist':'SHORTLISTED' if ats>=80 else 'CONSIDER','candidate':r,'job':j,'isSimulated':True}

def _safe_db_init():
 try: init_db()
 except Exception: pass

SIGNIN_HTML='''<!doctype html><html><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Sign in · IntelliHire</title><style>body{margin:0;min-height:100vh;display:grid;place-items:center;background:linear-gradient(135deg,#071426,#10274b 55%,#0b1230);font-family:system-ui,sans-serif;color:#fff}.card{width:min(420px,calc(100vw - 40px));padding:42px;border:1px solid rgba(255,255,255,.12);border-radius:24px;background:rgba(8,18,38,.86);text-align:center;box-shadow:0 24px 80px rgba(0,0,0,.35)}.brand{font-size:28px;font-weight:800}.eyebrow{color:#6fe7ff;font-size:12px;letter-spacing:.14em}.sub{color:#b7c6da;margin:10px 0 28px}.btn{display:block;padding:14px 18px;border-radius:12px;background:#fff;color:#132238;text-decoration:none;font-weight:700}</style></head><body><main class="card"><div class="eyebrow">AI-BASED PLACEMENT TRAINER</div><div class="brand">✦ INTELLIHIRE</div><h1>Welcome Back!</h1><div class="sub">Sign in to continue your placement journey.</div><a class="btn" href="/auth/google">Continue with Google</a></main></body></html>'''
DASH_HTML='''<!doctype html><html><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Dashboard · IntelliHire</title><style>body{margin:0;background:#071426;color:#fff;font-family:system-ui,sans-serif}.wrap{max-width:1100px;margin:auto;padding:36px}.top{display:flex;justify-content:space-between;align-items:center}.cards{display:grid;grid-template-columns:repeat(auto-fit,minmax(220px,1fr));gap:16px;margin-top:28px}.card{padding:24px;border:1px solid #203552;border-radius:18px;background:#0d1c33}.muted{color:#9db0c8}.btn{display:inline-block;margin:6px 8px 0 0;padding:12px 16px;border-radius:10px;background:#16345c;color:#fff;text-decoration:none}</style></head><body><div class="wrap"><div class="top"><div><div class="muted">INTELLIHIRE</div><h1>Welcome, {{ user.name or user.email }}</h1></div><a class="btn" href="/auth/logout">Sign out</a></div><div class="cards"><div class="card"><h2>Module 1</h2><p class="muted">ATS Resume Intelligence</p><a class="btn" href="/app/module1/overview">Open</a></div><div class="card"><h2>Module 2</h2><p class="muted">Adaptive Assessment</p><a class="btn" href="/app/module2/overview">Open</a></div><div class="card"><h2>Module 3</h2><p class="muted">Technical Interview</p><a class="btn" href="/app/module3">Open</a></div><div class="card"><h2>Module 4</h2><p class="muted">HR & Behavioral Interview</p><a class="btn" href="/app/module4">Open</a></div><div class="card"><h2>Module 5</h2><p class="muted">Hiring Committee & Readiness</p><a class="btn" href="/app/module5">Open</a></div></div></div></body></html>'''
ADMIN_HTML='''<!doctype html><html><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Admin · IntelliHire</title><style>body{margin:0;background:#071426;color:#fff;font-family:system-ui,sans-serif}.wrap{max-width:1250px;margin:auto;padding:28px}.top{display:flex;justify-content:space-between;align-items:center}.toolbar{display:flex;gap:10px;flex-wrap:wrap;margin:18px 0}.btn{padding:10px 14px;border-radius:9px;background:#173761;color:#fff;text-decoration:none;border:0}.table{width:100%;border-collapse:collapse;background:#0d1c33;border:1px solid #203552}.table th,.table td{padding:10px;border-bottom:1px solid #203552;text-align:left}.score{font-weight:700}.muted{color:#9db0c8}</style></head><body><div class="wrap"><div class="top"><div><div class="muted">INTELLIHIRE ADMIN</div><h1>Users & Performance</h1></div><a class="btn" href="/auth/logout">Sign out</a></div><form class="toolbar" method="get" action="/admin"><input name="q" value="{{ q or '' }}" placeholder="Search name or email" style="padding:10px;border-radius:9px;border:1px solid #35506f;background:#09182b;color:#fff"><button class="btn" type="submit">Search</button><a class="btn" href="/admin/reports.xlsx">Export All Excel</a></form><form method="post" action="/admin/reports.xlsx"><table class="table"><thead><tr><th></th><th>User</th><th>Email</th><th>Status</th><th>M1</th><th>M2</th><th>M3</th><th>M4</th><th>M5</th><th>Overall</th><th>Last Login</th></tr></thead><tbody>{% for u in users %}<tr><td><input type="checkbox" name="user_id" value="{{u.id}}"></td><td><a style="color:#8fe8ff" href="/admin/users/{{u.id}}">{{u.name or 'Unnamed'}}</a></td><td>{{u.email}}</td><td>{{u.status}}</td>{% for m in ['module1','module2','module3','module4','module5'] %}<td class="score">{{ u.scores[m].score if u.scores[m] else '—' }}</td>{% endfor %}<td class="score">{{u.overall if u.overall is not none else '—'}}</td><td>{{u.last_login_at or '—'}}</td></tr>{% endfor %}</tbody></table><div class="toolbar"><button class="btn" type="submit">Export Selected Excel</button></div></form></div></body></html>'''
ADMIN_USER_HTML='''<!doctype html><html><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Candidate · IntelliHire Admin</title><style>body{margin:0;background:#071426;color:#fff;font-family:system-ui,sans-serif}.wrap{max-width:1100px;margin:auto;padding:30px}.cards{display:grid;grid-template-columns:repeat(auto-fit,minmax(170px,1fr));gap:12px}.card{padding:18px;background:#0d1c33;border:1px solid #203552;border-radius:14px}.btn{display:inline-block;padding:10px 14px;border-radius:9px;background:#173761;color:#fff;text-decoration:none;margin:6px}.muted{color:#9db0c8}table{width:100%;border-collapse:collapse;background:#0d1c33}th,td{padding:9px;border-bottom:1px solid #203552;text-align:left}</style></head><body><div class="wrap"><a class="btn" href="/admin">← Users</a><h1>{{user.name or user.email}}</h1><p class="muted">{{user.email}} · {{user.status}}</p><div class="cards">{% for m in ['module1','module2','module3','module4','module5'] %}<div class="card"><div class="muted">{{m|upper}}</div><h2>{{scores[m].score if scores[m] else '—'}}</h2><a class="btn" href="/admin/users/{{user.id}}/modules/{{m}}">View</a></div>{% endfor %}</div><p><a class="btn" href="/admin/users/{{user.id}}/report.xlsx">Download Excel Report</a></p><h2>Activity</h2><table><tr><th>Time</th><th>Event</th><th>Metadata</th></tr>{% for a in activity %}<tr><td>{{a.created_at}}</td><td>{{a.event}}</td><td>{{a.metadata}}</td></tr>{% endfor %}</table></div></body></html>'''

@app.get('/')
def home():
    if current_user():
        user=current_user(); return redirect('/admin' if user.get('role')=='ADMIN' else '/app/dashboard')
    return render_template_string(SIGNIN_HTML)

@app.get('/app/dashboard')
def dashboard():
    user,response=require_auth('USER')
    if response:return response
    return render_template_string(DASH_HTML,user=user)

@app.get('/admin')
def admin_dashboard():
    user,response=require_auth('ADMIN')
    if response:return response
    _safe_db_init()
    try: users=list_user_summaries(request.args.get('q'))
    except Exception: users=[]
    return render_template_string(ADMIN_HTML,users=users,q=request.args.get('q',''))

@app.get('/admin/users/<user_id>')
def admin_user(user_id):
    user,response=require_auth('ADMIN')
    if response:return response
    _safe_db_init()
    users=[u for u in list_user_summaries() if u['id']==user_id]
    if not users: abort(404)
    from auth_db import get_user_activity
    return render_template_string(ADMIN_USER_HTML,user=users[0],scores=users[0]['scores'],activity=get_user_activity(user_id))

@app.get('/admin/users/<user_id>/modules/<module>')
def admin_user_module(user_id,module):
    user,response=require_auth('ADMIN')
    if response:return response
    if module not in {'module1','module2','module3','module4','module5'}: abort(404)
    _safe_db_init(); users=[u for u in list_user_summaries() if u['id']==user_id]
    if not users: abort(404)
    item=users[0]['scores'].get(module)
    return jsonify({'user':{'id':users[0]['id'],'name':users[0]['name'],'email':users[0]['email']},'module':module,'performance':item})

def _excel(user_ids):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    _safe_db_init()
    all_users=list_user_summaries()
    chosen=[u for u in all_users if not user_ids or u['id'] in set(user_ids)]
    wb=Workbook(); ws=wb.active; ws.title='Performance'
    headers=['User ID','Name','Email','Status','Module 1','Module 2','Module 3','Module 4','Module 5','Overall','Last Login']
    ws.append(headers)
    for c in ws[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='173761')
    for u in chosen:
        ws.append([u['id'],u['name'],u['email'],u['status'],*[u['scores'][m]['score'] if u['scores'][m] else None for m in ['module1','module2','module3','module4','module5']],u['overall'],u['last_login_at']])
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width=max(14,min(34,max(len(str(cell.value or '')) for cell in col)+2))
    activity_ws=wb.create_sheet('Activity')
    activity_ws.append(['User ID','Event','Timestamp','Metadata'])
    for c in activity_ws[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='173761')
    for u in chosen:
        for a in get_user_activity(u['id']): activity_ws.append([u['id'],a['event'],a['created_at'],json.dumps(a['metadata'])])
    output=BytesIO(); wb.save(output); output.seek(0); return output

@app.get('/admin/reports.xlsx')
def admin_report_all():
    user,response=require_auth('ADMIN')
    if response:return response
    return send_file(_excel([]),mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',as_attachment=True,download_name='intellihire_users_performance.xlsx')

@app.post('/admin/reports.xlsx')
def admin_report_selected():
    user,response=require_auth('ADMIN')
    if response:return response
    ids=request.form.getlist('user_id')
    if not ids: return redirect('/admin')
    return send_file(_excel(ids),mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',as_attachment=True,download_name='intellihire_selected_users.xlsx')

@app.get('/admin/users/<user_id>/report.xlsx')
def admin_report_user(user_id):
    user,response=require_auth('ADMIN')
    if response:return response
    return send_file(_excel([user_id]),mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',as_attachment=True,download_name=f'intellihire_{user_id}.xlsx')

@app.get('/app/<path:path>')
def shell(path):
    user=current_user()
    if not user:return redirect('/')
    if path.startswith('module2'): return render_template('module2.html')
    if path.startswith('module3'): return render_template('module3.html')
    if path.startswith('module4'): return render_template('module4.html')
    if path.startswith('module5'): return render_template('module5.html')
    if path.startswith('module1'): return render_template('index.html')
    return redirect('/app/dashboard')

@app.post('/app/module2/start')
def start_m2():
    user,response=require_auth('USER')
    if response:return response
    session['m2']={'ability':0.0,'answered':[],'events':[],'section':request.form.get('section',''),'count':int(request.form.get('count',8))};return redirect('/app/module2/assessment')
@app.route('/app/module2/assessment',methods=['GET','POST'])
def assessment():
    user,response=require_auth('USER')
    if response:return response
    s=session.get('m2',{'ability':0.0,'answered':[],'events':[],'section':'','count':8})
    if request.method=='POST':
        q=next((x for x in QUESTION_BANK if x.id==request.form.get('question_id')),None)
        if q:
            e=evaluate_answer(s['ability'],q,int(request.form.get('answer',-1)));s['events'].append(e);s['answered'].append(q.id);s['ability']=e['ability_after'];session['m2']=s
    if len(s['events'])>=s['count']:
        result=score_assessment(s['events'])
        try: record_performance(user['id'],'module2',result.get('score',result.get('overall',0)),result)
        except Exception: pass
        return redirect('/app/module2/results')
    q=select_next(s['ability'],s['answered'],s['section'] or None);return render_template('module2_assessment.html',question=q,ability=s['ability'],progress=len(s['events']),total=s['count'])
@app.get('/app/module2/results')
def m2_results():
    user,response=require_auth('USER')
    if response:return response
    result=score_assessment(session.get('m2',{}).get('events',[]))
    try: record_performance(user['id'],'module2',result.get('score',result.get('overall',0)),result)
    except Exception: pass
    return render_template('module2_results.html',result=result)

@app.get('/api/health')
def health():return jsonify({'status':'ok','engine':'IntelliHire Python ML Engine','version':'2.7','modules':['module1','module2','module3','module4','module5'],'auth':'enabled','admin_reporting':'enabled'})
@app.get('/api/demo')
def demo():return jsonify({'resume':DEMO,'job':JOB,'result':analyze({})})
@app.post('/api/analyze')
def api_analyze():
    user,response=require_auth('USER')
    if response:return response
    result=analyze(request.get_json(silent=True) or {})
    try: record_performance(user['id'],'module1',result.get('atsScore',0),result)
    except Exception: pass
    return jsonify(result)
if __name__=='__main__':app.run(host='0.0.0.0',port=int(os.getenv('PORT','5000')))
